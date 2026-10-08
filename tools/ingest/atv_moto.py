"""2026 ATV&MOTO.xlsx: несколько моделей на листе, цена EXW в юанях в колонке D."""
import io, re, sys, json, warnings
warnings.filterwarnings("ignore")
import openpyxl
from PIL import Image
from common import *

FILE = SRC + "/1iaHrmgYOMIkz9-Zx3SlNlVDZOLMTbUc4.xlsx"
SRC_NAME = "2026 ATV&MOTO.xlsx"

def kind(title):
    t = title.lower()
    if "kid" in t or "child" in t: age = "детский "
    else: age = ""
    if "off-road vehicle" in t or "off road vehicle" in t:
        return "moto", age, "мотоцикл (кросс/питбайк)"
    return "quad", age, "квадроцикл"

def main():
    wb = openpyxl.load_workbook(FILE)
    out = []
    for ws in wb.worksheets:
        cat, age, word = kind(ws.title)
        starts = []
        for r in range(1, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if a and re.match(r"^\s*[A-Za-z]{1,3}-[\w\-]+\s*$", str(a)):
                starts.append((r, str(a).strip()))
        starts.append((ws.max_row + 1, None))
        imgs = []
        for im in ws._images:
            try:
                data = im._data()
                pil = Image.open(io.BytesIO(data)); pil.load()
            except Exception:
                continue
            imgs.append((im.anchor._from.row + 1, pil.size[0] * pil.size[1], pil))
        for (r0, code), (r1, _) in zip(starts, starts[1:]):
            pairs, price, opts = [], None, []
            for r in range(r0, r1):
                lab, val = ws.cell(r, 2).value, ws.cell(r, 3).value
                if val and re.fullmatch(r"\s*\d[\d,]*\s*RMB\s*", str(val), re.I):
                    if price is None and r <= r0 + 3: price = int(re.sub(r"\D", "", str(val)))
                    val = None
                if lab and not val and ":" in str(lab):
                    lab, val = str(lab).split(":", 1)
                if lab and val and "model number" not in str(lab).lower():
                    pairs.append((lab, val))
                for c in (4, 5):
                    d = ws.cell(r, c).value
                    if not d: continue
                    d = str(d).strip()
                    m = re.fullmatch(r"\s*(\d[\d,]*)\s*RMB\s*", d, re.I)
                    if m and price is None and r <= r0 + 3:
                        price = int(m.group(1).replace(",", ""))
                    elif re.search(r"\+\s*\d+\s*rmb", d, re.I):
                        opts.append(re.sub(r"\s+", " ", d))
            mine = sorted([i for i in imgs if r0 - 1 <= i[0] < r1], key=lambda i: -i[1])
            pid = "xy-" + slug(code)
            paths = []
            for k, (_, area, pil) in enumerate(mine[:4]):
                pth = save_image(pil, f"atv-moto/{pid}-{k+1}")
                if pth and pth not in paths: paths.append(pth)
            p = {"id": pid, "cat": cat, "brand": "XY", "sku": code, "price": price, "cur": "RMB", "priceNote": "EXW", "images": paths, "src": SRC_NAME}
            title_cc = [v for l, v in pairs if re.search(r"engine|emission|power:", str(l).lower())]
            p["name"] = (word.capitalize() if not age else (age + word).capitalize()) + " " + code
            apply_specs(p, pairs)
            if "engine_cc" in p: p["name"] = f"{word.capitalize()} {int(p['engine_cc'])} см³ {age}".strip() + " " + code
            if opts: p["extra"].append(["Опции (доплата)", "; ".join(opts)[:400]])
            p["extra"] = [e for e in p["extra"] if e[0] != "Модель"]
            out.append(p)
    json.dump(out, open(OUT + "/atv_moto.json", "w"), ensure_ascii=False)
    print(len(out), "моделей;", sum(1 for p in out if p["price"]), "с ценой;", sum(1 for p in out if p["images"]), "с фото")

main()

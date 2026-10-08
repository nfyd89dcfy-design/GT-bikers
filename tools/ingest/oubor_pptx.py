"""欧铂尔产品图册.pptx (OUBOR): слайд = модель. Название в фигуре «OUBOR-…», характеристики в текстовом блоке, фото — картинки слайда."""
import io, re, json
from pptx import Presentation
from PIL import Image
from common import *

FID, FNAME = "1d9nC4cyVGlfB7QO9-ZlMs5Y__Ifq7pIW", "欧铂尔产品图册.pptx"
LAB = {"Length*Width*Height": "Габариты", "Lighting": "Освещение", "Tires Front": "Шины", "Tires": "Шины", "Brake F/R": "Тормоза (перед/зад)", "Brake": "Тормоза", "Motor": "Мотор",
       "BatteryCapacity": "Аккумулятор", "Controller": "Контроллер", "Seat Height": "Высота сиденья", "Ground Clearance": "Дорожный просвет", "Speed": "Скорость", "Range": "Запас хода",
       "Weight": "Масса", "Max Load": "Макс. нагрузка", "Charger": "Зарядное устройство", "Frame": "Рама", "Suspension": "Подвеска"}

def pairs_from(text):
    lines = [l.strip() for l in text.replace("\x0b", "\n").split("\n")]
    out, lab, buf = [], None, []
    for l in lines:
        if not l: continue
        if re.search(r"[:：]\s*$", l):
            if lab: out.append((lab, " ".join(buf)))
            lab, buf = re.sub(r"[:：]\s*$", "", l.split("/")[-1].strip()).strip(), []
            lab = re.sub(r"\s*\(.*?\)\s*", "", lab).strip()
        elif lab: buf.append(l)
    if lab: out.append((lab, " ".join(buf)))
    return [(a, b) for a, b in out if b]

def main():
    prs = Presentation(SRC + f"/{FID}.pptx"); out = []
    for si, s in enumerate(prs.slides, 1):
        name, spec, pics = None, None, []
        def walk(shapes):
            nonlocal name, spec
            for sh in shapes:
                if sh.shape_type == 6: walk(sh.shapes); continue
                if sh.has_text_frame:
                    t = sh.text_frame.text.strip()
                    if re.match(r"^OUBOR-", t) and not name: name = t.split("\n")[0].strip()
                    elif re.search(r"Lighting|Tires|Motor|Battery|Length", t): spec = (spec or "") + "\n" + t
                if sh.shape_type == 13:
                    if sh.top < 120 * 9525 and sh.height <= 125 * 9525: continue   # логотип
                    pics.append(sh)
        walk(s.shapes)
        if not name or not spec: continue
        code = name.replace("OUBOR-", "")
        pid = "oubor-" + slug(code)
        pairs = pairs_from(spec)
        p = {"id": pid, "cat": "ebike-city", "brand": "OUBOR", "name": "Электробайк OUBOR " + code, "sku": name, "images": [], "src": FNAME, "page": si, "extra": []}
        for a, b in pairs:
            ru = next((v for k, v in LAB.items() if a.replace(" ", "").lower().startswith(k.replace(" ", "").lower())), a)
            b = re.sub(r"\s+", " ", b).strip()
            if re.search(r"^(Габариты)", ru): b = b.replace("*", "×")
            p["extra"].append([ru, b])
            if ru == "Мотор":
                m = re.findall(r"(\d{3,5})\s*W", b, re.I)
                if m: p["power_kw"] = max(int(x) for x in m) / 1000
            if ru == "Аккумулятор":
                m = re.search(r"(\d{2,3})\s*V\s*(\d{1,3})\s*AH", b, re.I)
                if m:
                    p["battery_v"] = int(m.group(1)); p["battery_ah"] = int(m.group(2)); p["battery_wh"] = int(m.group(1)) * int(m.group(2))
            if ru == "Тормоза (перед/зад)" and re.search(r"disc", b, re.I): p["brakes"] = "Дисковые"
            if ru == "Тормоза (перед/зад)" and re.search(r"drum", b, re.I): p["brakes"] = "Барабанные"
        if p.get("power_kw", 0) >= 3: p["cat"] = "ebike-sport"
        pics.sort(key=lambda sh: -sh.width * sh.height)
        for k, sh in enumerate(pics[:4]):
            try:
                img = save_image(sh.image.blob, f"oubor/{pid}-{k+1}")
                if img and img not in p["images"]: p["images"].append(img)
            except Exception as e: print("img", si, e)
        out.append(p)
    json.dump(out, open(OUT + "/oubor_pptx.json", "w"), ensure_ascii=False)
    print(len(out), "моделей;", sum(1 for p in out if p["images"]), "с фото;", sum(1 for p in out if "power_kw" in p), "с мощностью;", sum(1 for p in out if "battery_wh" in p), "с АКБ")

main()

"""Прайс-листы экипировки и запчастей (FOX, 100%, USWE, Renthal, ProTaper, Polisport, Twin Air, Mobius).
Строка = артикул (цвет × размер). Объединяем строки в модель: размеры собираем в список."""
import io, re, json, sys, warnings
warnings.filterwarnings("ignore")
import openpyxl
from PIL import Image
from common import *

# (id файла, имя файла, бренд, категория по умолчанию)
FILES = [
 ("13m4-5EJBZrmG4UjLFGek3Iia3JZdl5Q5", "FOX MX Suit 202610.xlsx", "FOX", "gear"),
 ("1boBIFcdRi_I2wwa2lyGtdM2uogG_FV0p", "FOX MX Boot 202609.xlsx", "FOX", "gear"),
 ("1AF7uKCYDI8q-nqPGxhY89yh9A9HIMitc", "FOX goggle price list 202609.xlsx", "FOX", "gear"),
 ("19aYBICdpYq1gl_mGfb3fE-65L-F8a3wL", "FOX Glove 202609.xlsx", "FOX", "gear"),
 ("17QeIRHcJDDeJf_1n2aQpf8bQ0bGK3BNM", "FOX ADV price list.xlsx", "FOX", "gear"),
 ("1oF6sK5bMznM00kMsJplKY-DY1JWmWuYg", "FOX MTB clothing Price 202609.xlsx", "FOX", "gear"),
 ("1gRHl1xiT_rglV60puiMRTrBkWYUeAYyO", "USWE price list 202609.xlsx", "USWE", "gear"),
 ("1xkrHSl1QLXF6XF0Q86nLhi9f9LuryX5_", "100% Gloves.xlsx", "100%", "gear"),
 ("1gs0Rwkv0EXBXjtehwUuG6hyrmapFdSlm", "100% 202609 风镜报价单.xlsx", "100%", "gear"),
 ("1-sjfEJgqCHJEzhhNj9v_WZGqCYX_J10W", "Renthal 2026.xlsx", "Renthal", "parts"),
 ("1ljsCkT6L1OFjHvZBb2j7zNJtjKucJeI6", "Protaper 202609 price list.xlsx", "ProTaper", "parts"),
 ("1kFbU1LCBMJG8yMS-fueu7kXW5_8OIFda", "Polisport 260902.xlsx", "Polisport", "parts"),
 ("1lZBa3z6U2lHTCpCiP5TaBtM6p_mF1lqP", "Mobius price list.xlsx", "Mobius", "gear"),
]
GEAR_TYPES = [(r"goggle|风镜|lens|nose guard|tear|roll.?off", "Очки"), (r"glove|手套|mitt", "Перчатки"), (r"boot|靴", "Обувь"),
              (r"helmet|шлем|头盔", "Шлемы"), (r"jersey|jacket|shirt|tee|hoodie|ls |long sleeve|vest|fleece|短袖|衣", "Куртки и джерси"),
              (r"pant|short|bib|legging|裤", "Штаны и шорты"), (r"back ?pack|hydration|bag|pack|liner|bladder|包|水袋", "Рюкзаки и гидропаки"),
              (r"knee|brace|guard|protect|armou?r|elbow|护", "Защита")]
SIZES = {"XXS", "XS", "S", "M", "L", "XL", "2X", "2XL", "3X", "3XL", "XXL", "OS", "ONE SIZE", "SM", "MD", "LG", "S/M", "L/XL", "XS/S", "XL/2XL", "M/L"}
PARTS_TYPES = [(r"grip|ручк", "Ручки руля"), (r"bar|руль|handle", "Рули"), (r"sprocket|chain|звезд|cw|cs", "Звёзды и цепи"), (r"filter|фильтр|air|oil", "Фильтры"),
               (r"graphic|plastic|fender|shroud|guard|cover|fork|headlight|light|protect", "Пластик и защита")]
SERIES_RX = re.compile(r"^[^\d]*$")

def gtype(name, cat):
    n = name.lower()
    table = GEAR_TYPES if cat == "gear" else PARTS_TYPES
    for rx, t in table:
        if re.search(rx, n): return t
    return None

def num_(v):
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace(",", "").strip())
    except Exception: return None

def main():
    allp = []
    for fid, fname, brand, defcat in FILES:
        path = SRC + f"/{fid}.xlsx"
        wb = openpyxl.load_workbook(path, data_only=True)
        wf = openpyxl.load_workbook(path)           # формулы для =DISPIMG
        imgs = cell_images(path)
        for ws in wb.worksheets:
            wsf = wf[ws.title]
            # колонки цен по заголовку
            hdr = None
            for r in range(1, 8):
                row = {c.column: str(c.value).lower() for c in ws[r] if c.value}
                if any(re.search(r"retail|零售|新价格", v) for v in row.values()): hdr = (r, row); break
            if not hdr:
                for r in range(1, 8):
                    row = {c.column: str(c.value).lower() for c in ws[r] if c.value}
                    if any("wholesale" in v for v in row.values()): hdr = (r, row); break
            if not hdr: print("  нет заголовка:", fname, ws.title); continue
            hr, row = hdr
            retail = next((c for c, v in row.items() if re.search(r"retail|零售|新价格", v)), None)
            whole = next((c for c, v in row.items() if "wholesale" in v), None)
            if retail is None and whole: retail = whole - 1
            sizec = next((c for c, v in row.items() if v.strip() == "size"), None)
            fitc = next((c for c, v in row.items() if "fitment" in v), None)
            introc = next((c for c, v in row.items() if "介绍" in v), None)
            # floating images по строкам
            fl = []
            for im in getattr(wsf, "_images", []):
                try:
                    pil = Image.open(io.BytesIO(im._data())); pil.load(); fl.append((im.anchor._from.row + 1, pil))
                except Exception: pass
            series, groups, order = "", {}, []
            for r in range(hr + 1, ws.max_row + 1):
                cells = {c.column: c.value for c in ws[r] if c.value not in (None, "")}
                if not cells: continue
                rt = num_(cells.get(retail)) if retail else None
                texts = [(c, str(v).strip()) for c, v in sorted(cells.items()) if isinstance(v, str) and not str(v).startswith("=") and c not in (retail, whole)]
                if rt is None or rt <= 0:
                    if len(texts) == 1 and re.search(r"[A-Za-z一-鿿]", texts[0][1]) and texts[0][0] <= 3: series = texts[0][1]
                    continue
                # код и название
                texts = [t for t in texts if t[0] not in (sizec, fitc, introc)]
                if len(texts) < 2: continue
                code, name = texts[0][1], texts[1][1]
                if code.lower() in ("spec", "货号") or len(name) < 3: continue
                size = None
                if sizec and cells.get(sizec) is not None: size = str(cells[sizec]).strip()
                else:
                    suf = re.split(r"[-/]", code.split("/")[0])[-1].upper()
                    if suf in SIZES: size = suf
                wh = num_(cells.get(whole)) if whole else None
                key = (ws.title, name, str(cells.get(fitc)) if fitc else '')
                g = groups.get(key)
                if not g:
                    g = groups[key] = {"name": name, "codes": [], "sizes": [], "retail": rt, "whole": wh, "series": series, "row": r, "fit": None, "intro": None, "img": None, "sheet": ws.title}
                    order.append(key)
                g["codes"].append(code)
                if size and size not in g["sizes"]: g["sizes"].append(size)
                if fitc and cells.get(fitc) and not g["fit"]: g["fit"] = str(cells[fitc]).strip()
                if introc and cells.get(introc) and not g["intro"]: g["intro"] = str(cells[introc]).strip()
                if not g["img"]:
                    for cc in (1, 2, 3):
                        f = wsf.cell(r, cc).value
                        m = re.search(r'DISPIMG\("(ID_[^"]+)"', str(f)) if f else None
                        if m and m.group(1) in imgs: g["img"] = imgs[m.group(1)]; break
                    if not g["img"] and fl:
                        near = min(fl, key=lambda x: abs(x[0] - r))
                        if abs(near[0] - r) <= 2: g["img"] = near[1]
            n_img = 0
            for k in order:
                g = groups[k]
                pid = f"{slug(brand)}-{slug(g['codes'][0].split('/')[0])}"
                img = save_image(g["img"], f"gear/{pid}", max_side=640, quality=74, min_side=60) if g["img"] else None
                n_img += bool(img)
                gt = gtype(g["name"] + " " + ws.title + " " + (g["series"] or ""), defcat)
                p = {"id": pid, "cat": defcat, "brand": brand, "name": g["name"], "sku": g["codes"][0].split("/")[0], "price": g["retail"], "cur": "RMB", "priceNote": "розница",
                     "price2": g["whole"], "price2Note": "опт", "images": [img] if img else [], "src": f"{fname} · лист «{g['sheet']}»"}
                if defcat == "gear": p["gear_type"] = gt or "Экипировка"
                else: p["part_type"] = gt or "Запчасти и аксессуары"
                if g["sizes"]: p["sizes"] = g["sizes"]
                p["extra"] = []
                if g["series"]: p["extra"].append(["Серия", g["series"]])
                if g["fit"]: p["extra"].append(["Совместимость", g["fit"]])
                if g["intro"]: p["extra"].append(["Описание поставщика", g["intro"][:300]])
                if len(g["codes"]) > 1: p["extra"].append(["Артикулы", ", ".join(g["codes"][:12]) + (" …" if len(g["codes"]) > 12 else "")])
                allp.append(p)
            print(f"  {fname} / {ws.title}: моделей {len(order)}, с фото {n_img}")
    json.dump(allp, open(OUT + "/gear_xlsx.json", "w"), ensure_ascii=False)
    print("итого", len(allp))

main()
